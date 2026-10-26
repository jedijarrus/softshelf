"""
Excel-Export einer Mitteilung (v2.11).

Zwei Blaetter:
- Uebersicht: Kopfdaten + eine Zeile je Ausgabe (Zielgeraete, zugestellt, gesehen, ...)
- Empfaenger: eine Zeile je Ausgabe x Geraet x Windows-Benutzer, Zielgeraete ohne
  Rueckmeldung als "Nicht zugestellt"

build_workbook() ist synchron (openpyxl) und bekommt nur fertige Daten.
"""
from __future__ import annotations

import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SEVERITY_LABEL = {"info": "Info", "wichtig": "Wichtig", "kritisch": "Kritisch"}
STATE_LABEL = {
    "geplant": "Geplant", "aktiv": "Aktiv", "beendet": "Beendet",
    "pausiert": "Pausiert", "zurueckgezogen": "Zurückgezogen",
}
DATE_FMT = "DD.MM.YYYY HH:MM"

_HEAD_FONT = Font(bold=True)
_HEAD_FILL = PatternFill("solid", fgColor="F4F4F5")
_TITLE_FONT = Font(bold=True, size=13)
_KEY_FONT = Font(color="71717A")


def receipt_status(r: dict | None, critical: bool) -> str:
    if not r or not r.get("delivered_at"):
        return "Nicht zugestellt"
    if critical and r.get("acked_at"):
        return "Bestätigt"
    if r.get("shown_at") or r.get("acked_at"):
        return "Gesehen"
    return "Zugestellt"


def filename(message_id: int, title: str, today: datetime) -> str:
    t = title
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("Ä", "Ae"), ("Ö", "Oe"), ("Ü", "Ue"), ("ß", "ss")):
        t = t.replace(a, b)
    slug = re.sub(r"[^A-Za-z0-9]+", "-", t).strip("-")[:40] or "Mitteilung"
    return f"Mitteilung_{message_id}_{slug}_{today:%Y-%m-%d}.xlsx"


def _header_row(ws, row: int, labels: list[str]):
    for col, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=col, value=label)
        c.font = _HEAD_FONT
        c.fill = _HEAD_FILL


def _widths(ws, widths: list[int]):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(meta: list[tuple[str, object]], occurrences: list[dict],
                   rows: list[dict], critical: bool) -> bytes:
    """meta: (Bezeichnung, Wert)-Paare fuer den Kopf.
    occurrences: [{fire_at, expires_at, targets, delivered, shown, acked}] (naive datetimes)
    rows: [{fire_at, hostname, user, status, delivered_at, shown_at, acked_at, tray}]"""
    wb = Workbook()

    # ── Uebersicht ──
    ws = wb.active
    ws.title = "Übersicht"
    ws.cell(row=1, column=1, value="Mitteilung").font = _TITLE_FONT
    r = 3
    for key, value in meta:
        k = ws.cell(row=r, column=1, value=key)
        k.font = _KEY_FONT
        k.alignment = Alignment(vertical="top")
        c = ws.cell(row=r, column=2, value=value)
        # Wert ueber B:G, damit laengere Texte nicht die Tabellenspalten aufblaehen.
        # Excel passt die Hoehe verbundener Zellen nicht selbst an -> grob schaetzen.
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        if isinstance(value, datetime):
            c.number_format = DATE_FMT
        elif isinstance(value, str):
            lines = sum(max(1, -(-len(line) // 80)) for line in value.split("\n"))
            if lines > 1:
                ws.row_dimensions[r].height = 15 * min(lines, 30)
        r += 1
    r += 1
    labels = ["Ausgabe", "Ende", "Zielgeräte", "Zugestellt", "Gesehen"]
    if critical:
        labels.append("Bestätigt")
    labels.append("Ohne Rückmeldung")
    _header_row(ws, r, labels)
    head_row = r
    for o in occurrences:
        r += 1
        vals = [o["fire_at"], o["expires_at"], o["targets"], o["delivered"], o["shown"]]
        if critical:
            vals.append(o["acked"])
        vals.append(max(0, o["targets"] - o["delivered"]))
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=v)
            if isinstance(v, datetime):
                c.number_format = DATE_FMT
    if occurrences:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(labels))}{r}"
    _widths(ws, [18, 18, 12, 12, 12, 12, 18])

    # ── Empfaenger ──
    ws2 = wb.create_sheet("Empfänger")
    labels2 = ["Ausgabe", "Gerät", "Benutzer", "Status", "Zugestellt am", "Gesehen am"]
    if critical:
        labels2.append("Bestätigt am")
    labels2.append("Tray-Version")
    _header_row(ws2, 1, labels2)
    for i, row in enumerate(rows, 2):
        vals = [row["fire_at"], row["hostname"], row["user"], row["status"],
                row["delivered_at"], row["shown_at"]]
        if critical:
            vals.append(row["acked_at"])
        vals.append(row["tray"])
        for col, v in enumerate(vals, 1):
            c = ws2.cell(row=i, column=col, value=v)
            if isinstance(v, datetime):
                c.number_format = DATE_FMT
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(labels2))}{max(1, len(rows) + 1)}"
    _widths(ws2, [17, 20, 22, 17, 17, 17] + ([17] if critical else []) + [13])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
